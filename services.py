"""
services.py

Camada de serviços ("inteligência") do sistema de automação de controle
de frota de veículos da Administração Regional (GDF) - Decreto nº
47.091/2025.

Este módulo concentra as regras de negócio do sistema:

    - processar_csv_prime(caminho_csv):
        Lê e padroniza o relatório CSV mensal do portal "Prime
        Benefícios".

    - calcular_metricas_veiculo(df_veiculo, placa, km_acumulado_historico):
        Calcula as métricas de um veículo no mês (divisão em 4 semanas
        corridas + totais mensais + lista de abastecimentos) e monta o
        dicionário de contexto no formato exato exigido pelo Modelo 4 -
        Ficha Individual de Controle de Frota (Decreto nº 47.091/2025).

    - gerar_ficha_word(contexto, pasta_destino):
        Gera a ficha individual (.docx) de um veículo a partir do
        modelo modelos/modelo_ficha.docx, usando docxtpl e python-docx.

    - gerar_relatorio_consolidado_excel(lista_metricas, caminho_excel):
        Gera a planilha (.xlsx) com o resumo executivo da frota no mês.

Observação importante sobre o CSV do Prime Benefícios:
    Como o layout exato de colunas exportado pelo portal pode variar
    conforme o contrato/configuração, o mapeamento de nomes de coluna
    é feito através do dicionário MAPEAMENTO_COLUNAS_CSV (logo abaixo).
    Caso o arquivo real utilize nomes de coluna diferentes dos listados,
    basta ajustar esse dicionário — nenhuma outra parte do código
    precisa ser alterada.

Formato do contexto retornado por calcular_metricas_veiculo (Modelo 4):
    Seção I - Identificação:
        marca_modelo, placa, ano_fab_mod, combustivel, tipo_frota,
        responsavel, periodo_referencia.
    Seção II - Quilometragem por semana:
        s1_km_ini, s1_km_fim, s1_km_perc (... s2, s3, s4 ...),
        total_km_mensal, km_medio_acumulado.
    Seção III - Tabela dinâmica de abastecimentos:
        abastecimentos (lista de dicts com data, litros, valor, posto,
        km), total_litros, total_valor, media_kml.
"""

import calendar
import re
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from docxtpl import DocxTemplate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

# ---------------------------------------------------------------------------
# CONSTANTES DE CONFIGURAÇÃO DO MÓDULO
# ---------------------------------------------------------------------------

MAPEAMENTO_COLUNAS_CSV = {
    "Placa": "placa",
    "PLACA": "placa",
    "Data": "data",
    "Data Abastecimento": "data",
    "Data Transação": "data",
    "Data da Transação": "data",
    "Hora": "hora",
    "Hora Abastecimento": "hora",
    "Hora da Transação": "hora",
    "Litros": "litros",
    "Qtd Litros": "litros",
    "Quantidade (L)": "litros",
    "Quantidade Litros": "litros",
    "Qtde_Combustivel_Abastecido": "litros",
    "Valor Total": "valor_total",
    "Valor (R$)": "valor_total",
    "Valor Abastecimento": "valor_total",
    "Valor Transação": "valor_total",
    "Valor_Abastecimento": "valor_total",
    "Km Atual": "km_hodometro",
    "Hodômetro": "km_hodometro",
    "KM": "km_hodometro",
    "Km Rodado": "km_hodometro",
    "Quilometragem": "km_hodometro",
    "Posto": "posto",
    "Estabelecimento": "posto",
    "Local": "posto",
    "Nome Estabelecimento": "posto",
    "Nome_Posto": "posto",
}

NOMES_MESES_PT = {
    1: "Janeiro", 2: "Fevereiro", 3: "Março", 4: "Abril",
    5: "Maio", 6: "Junho", 7: "Julho", 8: "Agosto",
    9: "Setembro", 10: "Outubro", 11: "Novembro", 12: "Dezembro",
}

DADOS_COMPLEMENTARES_FROTA = {
    placa: {"combustivel": "Flex", "tipo_frota": "Oficial", "responsavel": "A definir"}
    for placa in config.CADASTRO_FROTA
}

FORMATO_KM_EXCEL = '#,##0.0'
FORMATO_LITROS_EXCEL = '#,##0.0'
FORMATO_MOEDA_EXCEL = '"R$" #,##0.00'
FORMATO_CONSUMO_EXCEL = '0.00'


# ---------------------------------------------------------------------------
# FUNÇÕES AUXILIARES INTERNAS
# ---------------------------------------------------------------------------

def _converter_para_float(valor: Any) -> float:
    """Converte valores diversos para float limpando símbolos e formatação BR."""
    if valor is None:
        return 0.0
    if isinstance(valor, (int, float)):
        return 0.0 if pd.isna(valor) else float(valor)

    texto = str(valor).strip()
    if not texto or texto.lower() in {"nan", "none", "-", "n/a", "na"}:
        return 0.0

    texto = re.sub(r"[^\d,.\-]", "", texto)
    texto = texto.replace(".", "").replace(",", ".")

    try:
        return float(texto)
    except ValueError:
        return 0.0


def _formatar_numero_br(valor: Optional[float], casas_decimais: int = 1) -> str:
    """Formata float no padrão brasileiro (separador de milhar . e decimal ,)."""
    if valor is None:
        return "N/A"
    texto = f"{valor:,.{casas_decimais}f}"
    texto = texto.replace(",", "@").replace(".", ",").replace("@", ".")
    return texto


def _formatar_km(valor: Optional[float]) -> str:
    return f"{_formatar_numero_br(valor, 1)} Km" if valor is not None else "N/A"


def _formatar_litros(valor: Optional[float]) -> str:
    return f"{_formatar_numero_br(valor, 1)} L" if valor is not None else "N/A"


def _formatar_moeda(valor: Optional[float]) -> str:
    return f"R$ {_formatar_numero_br(valor, 2)}" if valor is not None else "N/A"


def _formatar_consumo(valor: Optional[float]) -> str:
    return f"{_formatar_numero_br(valor, 2)} Km/L" if valor is not None else "N/A"


# ---------------------------------------------------------------------------
# 1. PROCESSAMENTO DO CSV (PRIME BENEFÍCIOS)
# ---------------------------------------------------------------------------

def processar_csv_prime(caminho_csv) -> pd.DataFrame:
    """Lê e padroniza o relatório CSV mensal exportado pelo portal Prime Benefícios."""
    caminho_csv = Path(caminho_csv)

    if not caminho_csv.exists():
        raise FileNotFoundError(f"Arquivo CSV não encontrado: {caminho_csv}")

    try:
        with open(caminho_csv, "r", encoding="utf-8-sig", errors="ignore") as f:
            primeira_linha = f.readline()
        pular_linhas = 1 if "SEP=" in primeira_linha.upper() else 0

        try:
            df_bruto = pd.read_csv(
                caminho_csv,
                sep=";",
                decimal=",",
                encoding="utf-8-sig",
                skiprows=pular_linhas,
            )
            if len(df_bruto.columns) <= 1:
                df_bruto = pd.read_csv(
                    caminho_csv,
                    sep=",",
                    encoding="utf-8-sig",
                    skiprows=pular_linhas,
                )
        except UnicodeDecodeError:
            df_bruto = pd.read_csv(
                caminho_csv,
                sep=";",
                decimal=",",
                encoding="latin1",
                skiprows=pular_linhas,
            )
            if len(df_bruto.columns) <= 1:
                df_bruto = pd.read_csv(
                    caminho_csv,
                    sep=",",
                    encoding="latin1",
                    skiprows=pular_linhas,
                )

    except pd.errors.EmptyDataError as erro:
        raise ValueError(f"O arquivo CSV está vazio: {caminho_csv}") from erro
    except pd.errors.ParserError as erro:
        raise ValueError(
            f"Falha ao interpretar o CSV (verifique se o separador é ';'): {erro}"
        ) from erro

    if df_bruto.empty:
        raise ValueError(f"O arquivo CSV não contém nenhuma linha de dados: {caminho_csv}")

    df_bruto.columns = [str(coluna).strip() for coluna in df_bruto.columns]

    mapeamento_presente = {
        coluna_original: coluna_padrao
        for coluna_original, coluna_padrao in MAPEAMENTO_COLUNAS_CSV.items()
        if coluna_original in df_bruto.columns
    }
    df = df_bruto.rename(columns=mapeamento_presente)

    colunas_obrigatorias = {"placa", "data", "litros", "valor_total", "km_hodometro"}
    colunas_faltantes = colunas_obrigatorias - set(df.columns)
    if colunas_faltantes:
        raise ValueError(
            f"O CSV não contém as colunas obrigatórias: {colunas_faltantes}. "
            f"Colunas encontradas: {list(df_bruto.columns)}."
        )

    df["placa"] = df["placa"].astype(str).str.strip().str.upper()

    for coluna in ("litros", "valor_total", "km_hodometro"):
        df[coluna] = df[coluna].apply(_converter_para_float)

    if "posto" in df.columns:
        df["posto"] = df["posto"].astype(str).str.strip()
        df["posto"] = df["posto"].replace({"nan": None, "": None, "None": None})
    else:
        df["posto"] = None

    if "hora" in df.columns:
        texto_data_hora = (
            df["data"].astype(str).str.strip() + " " + df["hora"].astype(str).str.strip()
        )
        df["data_hora"] = pd.to_datetime(texto_data_hora, dayfirst=True, errors="coerce")
    else:
        df["data_hora"] = pd.to_datetime(
            df["data"].astype(str).str.strip(), dayfirst=True, errors="coerce"
        )

    total_linhas_antes = len(df)
    df = df.dropna(subset=["data_hora"])
    df = df[df["placa"].notna() & (df["placa"] != "") & (df["placa"] != "NAN")]
    linhas_descartadas = total_linhas_antes - len(df)
    if linhas_descartadas > 0:
        print(
            f"[services] Aviso: {linhas_descartadas} linha(s) descartadas "
            f"por falta de placa/data válida."
        )

    if df.empty:
        raise ValueError("Após a limpeza, nenhuma linha válida restou no CSV.")

    df["data"] = df["data_hora"].dt.date
    df = df.sort_values(by=["placa", "data_hora"]).reset_index(drop=True)
    df = df[["placa", "data", "data_hora", "litros", "valor_total", "km_hodometro", "posto"]]

    print(
        f"[services] CSV processado com sucesso: {len(df)} transação(ões) "
        f"válida(s) de {df['placa'].nunique()} veículo(s)."
    )

    return df


# ---------------------------------------------------------------------------
# 2. CÁLCULO DE MÉTRICAS POR VEÍCULO
# ---------------------------------------------------------------------------

def calcular_metricas_veiculo(
    df_veiculo: pd.DataFrame,
    placa: str,
    km_acumulado_historico: Optional[float],
) -> Dict[str, Any]:
    """Calcula as métricas mensais de um veículo e monta o contexto da ficha."""
    placa = str(placa).strip().upper()
    df_completo = df_veiculo

    if not df_veiculo.empty and "placa" in df_veiculo.columns:
        df_veiculo = df_veiculo[
            df_veiculo["placa"].astype(str).str.strip().str.upper() == placa
        ].copy()

    if not df_veiculo.empty:
        coluna_ordenacao = "data_hora" if "data_hora" in df_veiculo.columns else "data"
        df_veiculo = df_veiculo.sort_values(by=coluna_ordenacao).reset_index(drop=True)

    dados_cadastrais = config.CADASTRO_FROTA.get(placa, {})
    marca = dados_cadastrais.get("marca", "N/A")
    modelo = dados_cadastrais.get("modelo", "N/A")
    marca_modelo = f"{marca} {modelo}".strip() if dados_cadastrais else "N/A"
    ano_fab_mod = dados_cadastrais.get("ano", "N/A")

    dados_complementares = DADOS_COMPLEMENTARES_FROTA.get(placa, {})
    combustivel = dados_complementares.get("combustivel", "N/A")
    tipo_frota = dados_complementares.get("tipo_frota", "N/A")
    responsavel = dados_complementares.get("responsavel", "A definir")

    if not df_veiculo.empty:
        primeira_data = df_veiculo["data"].iloc[0]
        if isinstance(primeira_data, pd.Timestamp):
            primeira_data = primeira_data.date()
        ano_referencia = primeira_data.year
        mes_referencia_num = primeira_data.month
    elif df_completo is not None and not df_completo.empty and "data" in df_completo.columns:
        periodos_do_arquivo = pd.to_datetime(df_completo["data"]).dt.to_period("M")
        periodo_inferido = periodos_do_arquivo.mode().iloc[0]
        ano_referencia = periodo_inferido.year
        mes_referencia_num = periodo_inferido.month
    else:
        hoje = date.today()
        ano_referencia = hoje.year
        mes_referencia_num = hoje.month

    nome_mes = NOMES_MESES_PT.get(mes_referencia_num, str(mes_referencia_num))
    periodo_referencia = f"{nome_mes} / {ano_referencia}"

    ultimo_dia_mes = calendar.monthrange(ano_referencia, mes_referencia_num)[1]
    limites_semanas = [(1, 7), (8, 14), (15, 21), (22, ultimo_dia_mes)]

    if not df_veiculo.empty:
        km_inicial_mes = float(df_veiculo["km_hodometro"].iloc[0])
        km_final_mes = float(df_veiculo["km_hodometro"].iloc[-1])
        dias_do_mes = pd.to_datetime(df_veiculo["data"]).dt.day
    else:
        km_inicial_mes = 0.0
        km_final_mes = 0.0
        dias_do_mes = pd.Series([], dtype="int64")

    semanas_calculadas = []
    km_inicial_semana_atual = km_inicial_mes

    for dia_inicio, dia_fim in limites_semanas:
        if not df_veiculo.empty:
            mascara_semana = (dias_do_mes >= dia_inicio) & (dias_do_mes <= dia_fim)
            transacoes_semana = df_veiculo[mascara_semana]
        else:
            transacoes_semana = df_veiculo

        if not transacoes_semana.empty:
            km_final_semana = float(transacoes_semana["km_hodometro"].iloc[-1])
        else:
            km_final_semana = km_inicial_semana_atual

        km_percorrido_semana = km_final_semana - km_inicial_semana_atual
        semanas_calculadas.append((km_inicial_semana_atual, km_final_semana, km_percorrido_semana))
        km_inicial_semana_atual = km_final_semana

    (s1_ini, s1_fim, s1_perc), (s2_ini, s2_fim, s2_perc), \
        (s3_ini, s3_fim, s3_perc), (s4_ini, s4_fim, s4_perc) = semanas_calculadas

    total_km_mensal = round(km_final_mes - km_inicial_mes, 1)
    total_litros = round(float(df_veiculo["litros"].sum()), 2) if not df_veiculo.empty else 0.0
    total_valor = round(float(df_veiculo["valor_total"].sum()), 2) if not df_veiculo.empty else 0.0
    media_kml = round(total_km_mensal / total_litros, 2) if total_litros > 0 else 0.0

    abastecimentos: List[Dict[str, str]] = []
    for _, linha in df_veiculo.iterrows():
        data_abastecimento = linha["data"]
        if isinstance(data_abastecimento, pd.Timestamp):
            data_abastecimento = data_abastecimento.date()

        posto_valor = linha.get("posto")
        if pd.isna(posto_valor) or not str(posto_valor).strip():
            posto_valor = "Não informado"

        abastecimentos.append({
            "data": data_abastecimento.strftime("%d/%m/%Y"),
            "litros": _formatar_litros(float(linha["litros"])),
            "valor": _formatar_moeda(float(linha["valor_total"])),
            "posto": str(posto_valor),
            "km": _formatar_km(float(linha["km_hodometro"])),
        })

    contexto: Dict[str, Any] = {
        "marca_modelo": marca_modelo,
        "placa": placa,
        "ano_fab_mod": ano_fab_mod,
        "combustivel": combustivel,
        "tipo_frota": tipo_frota,
        "responsavel": responsavel,
        "periodo_referencia": periodo_referencia,

        "s1_km_ini": _formatar_km(s1_ini), "s1_km_fim": _formatar_km(s1_fim), "s1_km_perc": _formatar_km(s1_perc),
        "s2_km_ini": _formatar_km(s2_ini), "s2_km_fim": _formatar_km(s2_fim), "s2_km_perc": _formatar_km(s2_perc),
        "s3_km_ini": _formatar_km(s3_ini), "s3_km_fim": _formatar_km(s3_fim), "s3_km_perc": _formatar_km(s3_perc),
        "s4_km_ini": _formatar_km(s4_ini), "s4_km_fim": _formatar_km(s4_fim), "s4_km_perc": _formatar_km(s4_perc),
        "total_km_mensal": _formatar_km(total_km_mensal),
        "km_medio_acumulado": (
            _formatar_consumo(km_acumulado_historico)
            if km_acumulado_historico is not None
            else "Sem histórico anterior"
        ),

        "abastecimentos": abastecimentos,
        "total_litros": _formatar_litros(total_litros),
        "total_valor": _formatar_moeda(total_valor),
        "media_kml": _formatar_consumo(media_kml),

        "mes_numero": f"{mes_referencia_num:02d}",
        "ano_referencia": str(ano_referencia),
        "data_emissao": date.today().strftime("%d/%m/%Y"),
        "total_km_mensal_num": total_km_mensal,
        "total_litros_num": total_litros,
        "total_valor_num": total_valor,
        "media_kml_num": media_kml,
        "km_medio_acumulado_num": km_acumulado_historico,
    }

    return contexto


# ---------------------------------------------------------------------------
# 3. GERAÇÃO DA FICHA INDIVIDUAL EM WORD
# ---------------------------------------------------------------------------

def gerar_ficha_word(contexto: Dict[str, Any], pasta_destino) -> Path:
    """
    Gera a ficha individual de controle em formato Word (.docx).
    Usa docxtpl para renderizar os campos de cabeçalho e python-docx para
    preencher a tabela de abastecimentos dinamicamente sem erros de tags Jinja.
    """
    if "placa" not in contexto:
        raise KeyError("O dicionário de contexto precisa conter a chave 'placa'.")

    caminho_modelo = config.PASTA_MODELOS / "modelo_ficha.docx"
    if not caminho_modelo.exists():
        raise FileNotFoundError(f"Modelo de ficha não encontrado em: {caminho_modelo}.")

    pasta_destino = Path(pasta_destino)

    try:
        pasta_destino.mkdir(parents=True, exist_ok=True)

        documento = DocxTemplate(str(caminho_modelo))

        # 1. Renderiza os campos fixos do cabeçalho e totais
        documento.render(contexto)

       # 2. Preenchimento direto da Tabela III (Abastecimentos) via python-docx
        if len(documento.docx.tables) >= 3:
            tabela_consumo = documento.docx.tables[2]
            abastecimentos = contexto.get("abastecimentos", [])

            for item in abastecimentos:
                # Garante que a nova linha seja inserida ANTES da linha de totais
                linha_totais = tabela_consumo.rows[-1]._tr
                nova_linha = tabela_consumo.add_row()
                linha_totais.addprevious(nova_linha._tr)

                row_cells = nova_linha.cells
                if len(row_cells) >= 5:
                    row_cells[0].text = str(item.get("data", ""))
                    row_cells[1].text = str(item.get("litros", ""))
                    row_cells[2].text = str(item.get("valor", ""))
                    row_cells[3].text = str(item.get("posto", ""))
                    row_cells[4].text = str(item.get("km", ""))

        # 3. Salva o documento gerado
        placa = str(contexto["placa"]).strip().upper()
        mes_numero = contexto.get("mes_numero", "00")
        ano_referencia = contexto.get("ano_referencia", "0000")
        mes_ano_arquivo = f"{mes_numero}-{ano_referencia}"

        nome_arquivo = f"Ficha_CONTROLE_{placa}_{mes_ano_arquivo}.docx"
        caminho_saida = pasta_destino / nome_arquivo

        documento.save(str(caminho_saida))

        print(f"[services] Ficha gerada com sucesso: {caminho_saida}")
        return caminho_saida

    except (KeyError, FileNotFoundError):
        raise
    except Exception as erro:
        print(
            f"[services] ERRO ao gerar ficha Word para a placa "
            f"{contexto.get('placa', '???')}: {erro}"
        )
        raise


# ---------------------------------------------------------------------------
# 4. GERAÇÃO DO RELATÓRIO CONSOLIDADO EM EXCEL
# ---------------------------------------------------------------------------

def gerar_relatorio_consolidado_excel(lista_metricas: List[Dict[str, Any]], caminho_excel) -> Path:
    """Gera a planilha Excel consolidada com o resumo executivo da frota no mês."""
    caminho_excel = Path(caminho_excel)

    if not lista_metricas:
        print("[services] Aviso: lista_metricas está vazia.")

    try:
        caminho_excel.parent.mkdir(parents=True, exist_ok=True)

        workbook = Workbook()
        planilha = workbook.active
        planilha.title = "Resumo Frota"

        cabecalhos = [
            "Placa", "Marca/Modelo", "Km Percorrido",
            "Litros Totais", "Valor Total (R$)", "Média (Km/L)",
        ]
        planilha.append(cabecalhos)

        fonte_cabecalho = Font(bold=True, color="FFFFFF", size=11)
        preenchimento_cabecalho = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        alinhamento_cabecalho = Alignment(horizontal="center", vertical="center")

        for celula in planilha[1]:
            celula.font = fonte_cabecalho
            celula.fill = preenchimento_cabecalho
            celula.alignment = alinhamento_cabecalho

        for metricas in lista_metricas:
            placa = metricas.get("placa", "N/A")
            marca_modelo = metricas.get("marca_modelo", "N/A") or "N/A"
            km_percorrido = metricas.get("total_km_mensal_num", 0.0) or 0.0
            litros_totais = metricas.get("total_litros_num", 0.0) or 0.0
            valor_total = metricas.get("total_valor_num", 0.0) or 0.0
            media_kml = metricas.get("media_kml_num", 0.0) or 0.0

            planilha.append([placa, marca_modelo, km_percorrido, litros_totais, valor_total, media_kml])

            linha_atual = planilha.max_row
            planilha.cell(row=linha_atual, column=3).number_format = FORMATO_KM_EXCEL
            planilha.cell(row=linha_atual, column=4).number_format = FORMATO_LITROS_EXCEL
            planilha.cell(row=linha_atual, column=5).number_format = FORMATO_MOEDA_EXCEL
            planilha.cell(row=linha_atual, column=6).number_format = FORMATO_CONSUMO_EXCEL

        for indice_coluna, cabecalho in enumerate(cabecalhos, start=1):
            letra_coluna = get_column_letter(indice_coluna)
            maior_largura = len(cabecalho)
            for linha in planilha.iter_rows(min_col=indice_coluna, max_col=indice_coluna, min_row=2):
                for celula in linha:
                    if celula.value is not None:
                        maior_largura = max(maior_largura, len(str(celula.value)))
            planilha.column_dimensions[letra_coluna].width = maior_largura + 4

        planilha.freeze_panes = "A2"
        workbook.save(str(caminho_excel))

        print(
            f"[services] Planilha consolidada gerada com sucesso: {caminho_excel} "
            f"({len(lista_metricas)} veículo(s))."
        )
        return caminho_excel

    except Exception as erro:
        print(f"[services] ERRO ao gerar planilha consolidada: {erro}")
        raise